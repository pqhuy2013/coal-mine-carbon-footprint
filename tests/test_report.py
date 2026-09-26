from io import BytesIO

import pytest
from openpyxl import load_workbook

from coal_mine_footprint.combustion import CombustionType, Fuel
from coal_mine_footprint.inventory import FuelUse, InventoryInput, compute_inventory
from coal_mine_footprint.report import build_excel_report


def test_excel_report_contents():
    inp = InventoryInput(
        facility_name="Mỏ than A",
        reporting_year=2025,
        coal_production_t=1_000_000,
        fuels=[FuelUse(Fuel.DIESEL, CombustionType.MOBILE, 1000)],
        electricity_kwh=1_000_000,
        grid_ef_t_per_mwh=0.6,
    )
    result = compute_inventory(inp)
    wb = load_workbook(BytesIO(build_excel_report(inp, result)))

    assert wb.sheetnames == ["Thông tin chung", "Kết quả", "Số liệu đầu vào", "Hệ số phát thải"]
    assert wb["Thông tin chung"]["B3"].value == "Mỏ than A"

    ws = wb["Kết quả"]
    values = {row[1]: row[5] for row in ws.iter_rows(min_row=2, values_only=True) if row[1]}
    assert values["TỔNG PHÁT THẢI (tCO2e)"] == pytest.approx(result.total_co2e_t)
    assert values["Tổng Scope 2 (tCO2e)"] == pytest.approx(600)
    assert values["Cường độ phát thải (kgCO2e/tấn than)"] == pytest.approx(
        result.intensity_kg_per_t
    )

    factors = [row[0] for row in wb["Hệ số phát thải"].iter_rows(min_row=2, values_only=True)]
    assert "Dầu diesel (nguồn di động): CO2" in factors
