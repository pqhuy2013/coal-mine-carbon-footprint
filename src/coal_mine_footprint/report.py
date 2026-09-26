"""Xuất kết quả kiểm kê KNK ra tệp Excel."""

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from coal_mine_footprint.combustion import FUELS, CombustionType, Fuel, fuel_use_label
from coal_mine_footprint.ghg import GWP100
from coal_mine_footprint.inventory import (
    InventoryInput,
    InventoryResult,
    MethaneTier,
    methane_basis,
)
from coal_mine_footprint.methane import CH4_DENSITY_KG_PER_M3

NUMBER_FORMAT = "#,##0.00"
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="2F5597")
BOLD = Font(bold=True)

IPCC_COAL = "IPCC 2006, Tập 2, Chương 4.1"
IPCC_FUEL = "IPCC 2006, Tập 2, Chương 1-3"


def _header(ws: Worksheet, row: int, titles: list[str]) -> None:
    for col, title in enumerate(titles, start=1):
        cell = ws.cell(row=row, column=col, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _widths(ws: Worksheet, widths: list[int]) -> None:
    for i, width in enumerate(widths):
        ws.column_dimensions[chr(ord("A") + i)].width = width


def _info_sheet(ws: Worksheet, inp: InventoryInput) -> None:
    ws.title = "Thông tin chung"
    ws["A1"] = "BÁO CÁO KIỂM KÊ KHÍ NHÀ KÍNH - MỎ THAN HẦM LÒ"
    ws["A1"].font = Font(bold=True, size=14)
    rows = [
        ("Tên cơ sở", inp.facility_name),
        ("Địa chỉ", inp.address),
        ("Năm kiểm kê", inp.reporting_year),
        ("Người lập", inp.prepared_by),
        ("Ngày lập", date.today().strftime("%d/%m/%Y")),
        ("Sản lượng than nguyên khai (tấn)", inp.coal_production_t),
        ("Phương pháp", "IPCC 2006 Guidelines; GHG Protocol Corporate Standard"),
        ("Bộ hệ số GWP", f"IPCC {inp.gwp} (100 năm)"),
        (
            "Lưu ý",
            "Cần đối chiếu và chuyển số liệu sang biểu mẫu báo cáo theo quy định hiện hành.",
        ),
    ]
    for i, (label, value) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=label).font = BOLD
        ws.cell(row=i, column=2, value=value)
    ws["B8"].number_format = NUMBER_FORMAT
    _widths(ws, [34, 80])


def _results_sheet(ws: Worksheet, result: InventoryResult) -> None:
    _header(ws, 1, ["Phạm vi", "Nguồn phát thải", "CO2 (tấn)", "CH4 (tấn)", "N2O (tấn)", "tCO2e"])
    row = 2
    for line in result.lines:
        g = line.gases
        values = [f"Scope {line.scope}", line.source, g.co2_t, g.ch4_t, g.n2o_t]
        for col, value in enumerate(values + [g.co2e_t(result.gwp)], start=1):
            ws.cell(row=row, column=col, value=value)
        row += 1

    row += 1
    summary = [
        ("Tổng Scope 1 (tCO2e)", result.scope_co2e_t(1)),
        ("Tổng Scope 2 (tCO2e)", result.scope_co2e_t(2)),
        ("TỔNG PHÁT THẢI (tCO2e)", result.total_co2e_t),
        ("Cường độ phát thải (kgCO2e/tấn than)", result.intensity_kg_per_t),
    ]
    for label, value in summary:
        ws.cell(row=row, column=2, value=label).font = BOLD
        ws.cell(row=row, column=6, value=value).font = BOLD
        row += 1

    for cells in ws.iter_rows(min_row=2, min_col=3, max_col=6):
        for cell in cells:
            cell.number_format = NUMBER_FORMAT
    _widths(ws, [10, 44, 14, 14, 14, 16])


def _inputs_sheet(ws: Worksheet, inp: InventoryInput) -> None:
    _header(ws, 1, ["Hạng mục", "Giá trị", "Đơn vị"])
    tier = MethaneTier(inp.methane_tier)
    rows: list[tuple[str, object, str]] = [
        ("Sản lượng than nguyên khai", inp.coal_production_t, "tấn"),
        ("Cấp độ tính CH4", f"Tier {tier.value}", ""),
    ]
    if tier is MethaneTier.TIER3:
        rows.append(("Thể tích CH4 đo được trong khai thác", inp.measured_mining_m3, "m³"))
    rows.append(("CH4 thu hồi và đốt tại mỏ", inp.recovered_combusted_m3, "m³"))
    for use in inp.fuels:
        unit = FUELS[Fuel(use.fuel)].unit
        rows.append((fuel_use_label(use.fuel, use.combustion), use.quantity, unit))
    rows += [
        ("Điện năng mua từ lưới", inp.electricity_kwh, "kWh"),
        ("Khối lượng thuốc nổ", inp.explosives_kg, "kg"),
    ]
    for i, row in enumerate(rows, start=2):
        for col, value in enumerate(row, start=1):
            ws.cell(row=i, column=col, value=value)
        ws.cell(row=i, column=2).number_format = NUMBER_FORMAT
    _widths(ws, [44, 18, 10])


def _factors_sheet(ws: Worksheet, inp: InventoryInput) -> None:
    _header(ws, 1, ["Hệ số", "Giá trị", "Đơn vị", "Nguồn"])
    basis = methane_basis(inp)
    tier = MethaneTier(inp.methane_tier)
    methane_source = IPCC_COAL if tier is MethaneTier.TIER1 else "Số liệu của cơ sở"
    rows: list[tuple[str, object, str, str]] = []
    if basis.mining_ef is not None:
        rows.append(("Hệ số CH4 khai thác hầm lò", basis.mining_ef, "m³/tấn", methane_source))
    post_source = IPCC_COAL if inp.post_mining_ef is None else "Số liệu của cơ sở"
    rows += [
        ("Hệ số CH4 sau khai thác", basis.post_mining_ef, "m³/tấn", post_source),
        ("Khối lượng riêng CH4", CH4_DENSITY_KG_PER_M3, "kg/m³", IPCC_COAL),
    ]
    for key in sorted({(Fuel(u.fuel), CombustionType(u.combustion)) for u in inp.fuels}):
        fuel, combustion = key
        props = FUELS[fuel]
        ch4, n2o = (
            props.stationary_kg_per_tj
            if combustion is CombustionType.STATIONARY
            else props.mobile_kg_per_tj or (0.0, 0.0)
        )
        name = fuel_use_label(fuel, combustion)
        if fuel in inp.densities_kg_per_l:
            rows.append((f"{name}: khối lượng riêng", inp.densities_kg_per_l[fuel], "kg/lít", ""))
        rows += [
            (f"{name}: nhiệt trị thực", props.ncv_tj_per_gg, "TJ/Gg", IPCC_FUEL),
            (f"{name}: CO2", props.co2_kg_per_tj, "kg/TJ", IPCC_FUEL),
            (f"{name}: CH4", ch4, "kg/TJ", IPCC_FUEL),
            (f"{name}: N2O", n2o, "kg/TJ", IPCC_FUEL),
        ]
    gwp = GWP100[inp.gwp]
    rows += [
        ("Hệ số phát thải thuốc nổ", inp.explosives_ef_t_per_t, "tCO2/tấn", "Người dùng nhập"),
        ("Hệ số phát thải lưới điện", inp.grid_ef_t_per_mwh, "tCO2/MWh", "Người dùng nhập"),
        ("GWP CH4", gwp["CH4"], "", f"IPCC {inp.gwp}"),
        ("GWP N2O", gwp["N2O"], "", f"IPCC {inp.gwp}"),
    ]
    for i, row in enumerate(rows, start=2):
        for col, value in enumerate(row, start=1):
            ws.cell(row=i, column=col, value=value)
    _widths(ws, [48, 14, 12, 32])


def build_excel_report(inp: InventoryInput, result: InventoryResult) -> bytes:
    """Tạo báo cáo Excel gồm thông tin chung, kết quả, số liệu đầu vào và hệ số."""
    wb = Workbook()
    _info_sheet(wb.active, inp)
    _results_sheet(wb.create_sheet("Kết quả"), result)
    _inputs_sheet(wb.create_sheet("Số liệu đầu vào"), inp)
    _factors_sheet(wb.create_sheet("Hệ số phát thải"), inp)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
